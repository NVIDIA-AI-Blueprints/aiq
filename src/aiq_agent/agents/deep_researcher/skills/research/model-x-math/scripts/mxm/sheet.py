"""
ModelSheet — the named-row registry and layout engine.

Two phases so cross-block formulas work regardless of visual order:

  1. RESERVE — blocks call section()/line()/scalar()/spacer(). Each call assigns
     the next row number and records a render spec. Nothing is written yet, but
     `self.rows[name]` is populated immediately.
  2. RENDER — render() walks the specs and writes cells. Formula callables are
     evaluated here, with the FULL registry available, so a P&L "Interest" row
     can reference the Debt Schedule that sits below it (forward reference).

Formula callables have the signature `fn(S, i, col) -> str | None`:
  S   = this ModelSheet (use S.rref(name)/S.cref(name, i) to reference rows),
  i   = period index, col = that period's column letter. Return None to skip.
"""

from __future__ import annotations

from collections.abc import Callable
from collections.abc import Sequence

from openpyxl.comments import Comment
from openpyxl.worksheet.worksheet import Worksheet

from . import style
from .periods import PeriodAxis

Fn = Callable[["ModelSheet", int, str], str | None]


class ModelSheet:
    def __init__(self, ws: Worksheet, axis: PeriodAxis, company: str, units: str, title: str | None = None):
        self.ws = ws
        self.axis = axis
        self.company = company
        self.units = units
        self.title = title or f"{company} - Model X Math"
        self.rows: dict[str, int] = {}
        self._specs: list[dict] = []
        self.r = 14  # rows 1-13 are reserved for the fixed header (drawn in render)
        self._fx_row_name: str | None = None  # currency-conversion FX row (if any)
        self._fx_curr_name: str | None = None  # currency toggle cell (1 = source, 2 = reporting)
        self._fx_op = "/"
        ws.sheet_view.showGridLines = False

    # ---- references (used inside formula callables) -------------------------
    def rref(self, name: str) -> int:
        return self.rows[name]

    def set_fx(self, row_name: str, curr_name: str | None = None, op: str = "/") -> None:
        """Enable currency conversion as an in-Excel toggle. Every money input becomes
        `=<native>/CHOOSE($G$<curr>, 1, <this-year's FX cell>)`: with the toggle at 1 the
        cell shows the ORIGINAL source-currency number, at 2 it converts by that year's rate.
        Editing the toggle or any FX rate re-drives the whole model. `row_name` is the FX-rate
        row; `curr_name` is the toggle cell (both registered lines)."""
        self._fx_row_name = row_name
        self._fx_curr_name = curr_name
        self._fx_op = op

    def scenario_cell(self, names: Sequence[str], default: int = 1) -> int:
        """Reserve the scenario selector: a blue editable cell (col G) holding 1..N,
        registered as `scenario` (→ `$G$<row>`) for `CHOOSE($G$scen, caseA, caseB)`
        forecast toggles. `names` = the ordered case labels shown in the legend."""
        legend = ", ".join(f"{i + 1} = {n}" for i, n in enumerate(names))
        self.section("Scenario")
        row = self.line("scenario", label=f"Case ({legend})", g=float(default), g_fmt="0", currency=False)
        self.spacer(2)
        return row

    def cref(self, name: str, i: int) -> str:
        return f"{self.axis.col(i)}{self.rows[name]}"

    def xref(self, name: str, col: str) -> str:
        """Cross-sheet reference to THIS sheet's registered row `name` at column `col`,
        e.g. `'Company - Model X Math'!O45` — used by the DCF tab."""
        return f"'{self.ws.title}'!{col}{self.rows[name]}"

    # ---- reservation --------------------------------------------------------
    def spacer(self, n: int = 1) -> None:
        self.r += n

    def _reserve(self, spec: dict, name: str | None) -> int:
        row = self.r
        spec["row"] = row
        if name:
            self.rows[name] = row
        self._specs.append(spec)
        self.r += 1
        return row

    def section(self, title: str, *, years: bool = False, name: str | None = None) -> int:
        return self._reserve({"type": "section", "title": title, "years": years}, name)

    def line(
        self,
        name: str | None = None,
        *,
        label: str | None = None,
        fn: Fn | None = None,
        values: Sequence | None = None,
        fmt: str = style.NUM1,
        bold: bool = False,
        italic: bool = False,
        total: bool = False,
        assumption: bool = False,
        summary: bool = False,
        cols: Sequence[int] | None = None,
        g: float | None = None,
        g_fmt: str | None = None,
        g_fx_col: int = 0,
        currency: bool = True,
        comment: str | None = None,
    ) -> int:
        """A data row. Use `values` for blue inputs, `fn` for black formulas.

        `g` drops a single blue assumption input in col G on this row (as in the reference models:
        Tax Rate %, Interest %, Capex split), referenceable as `$G${row}`.

        `comment` attaches a cell note (the assumption's "why + basis") to the col-G input
        when `g` is set, else to the row label — used by the inline Assumptions block.

        `currency=True` (default) marks the row's `values` as absolute-currency amounts,
        so they are FX-converted when `set_fx()` is active. Set `currency=False` for
        rows whose `values` are not money (percentages, multiples, and the FX row itself).
        """
        return self._reserve(
            {
                "type": "line",
                "name": name,
                "label": name if label is None else label,
                "fn": fn,
                "values": list(values) if values is not None else None,
                "fmt": fmt,
                "bold": bold,
                "italic": italic,
                "total": total,
                "assumption": assumption,
                "summary": summary,
                "cols": list(cols) if cols is not None else None,
                "g": g,
                "g_fmt": g_fmt or fmt,
                "g_fx_col": g_fx_col,
                "currency": currency,
                "comment": comment,
            },
            name,
        )

    def sub(
        self, name: str | None, fn: Fn, *, label: str, fmt: str = style.PCT, cols: Sequence[int] | None = None
    ) -> int:
        """An italic %-sub-row (% Growth, % Margin, % of Revenue)."""
        return self.line(name, label=label, fn=fn, fmt=fmt, italic=True, cols=cols)

    # convenience derived rows -------------------------------------------------
    def growth(self, of: str, *, label: str = "% Growth") -> int:
        return self.sub(
            None,
            lambda S, i, c: f"{c}{S.rref(of)}/{S.axis.col(i - 1)}{S.rref(of)}-1",
            label=label,
            cols=range(1, self.axis.n),
        )

    def margin(
        self, of: str, *, base: str = "Revenue", label: str = "% Margin", cols: Sequence[int] | None = None
    ) -> int:
        return self.sub(None, lambda S, i, c: f"{c}{S.rref(of)}/{c}${S.rref(base)}", label=label, cols=cols)

    def check(self, label: str, a: str, b: str) -> int:
        return self.line(
            None, label=label, italic=True, fmt=style.NUM1, fn=lambda S, i, c: f"{c}{S.rref(a)}-{c}{S.rref(b)}"
        )

    # ---- render -------------------------------------------------------------
    def render(self) -> None:
        self._render_header()
        for s in self._specs:
            getattr(self, f"_r_{s['type']}")(s)
        self._widths()

    def _cell(self, col_idx: int, row: int):
        return self.ws.cell(row=row, column=col_idx)

    def _render_header(self) -> None:
        ax = self.axis
        c = self._cell(style.COL_LABEL, 2)
        c.value = self.title
        style.title(c)
        for lbl, row, kind in [("Start date", 6, "start"), ("End date", 7, "end"), ("Days per year", 9, "days")]:
            lc = self._cell(style.COL_LABEL, row)
            lc.value = lbl
            style.header_label(lc)
            for i in range(ax.n):
                cell = self._cell(ax.col_idx(i), row)
                if kind == "days":
                    cell.value = ax.days(i)
                    style.value(cell, fmt="0")
                else:
                    cell.value = ax.start_date(i) if kind == "start" else ax.end_date(i)
                    style.value(cell, fmt=style.DATE)
        # band labels row 11
        hist = self._cell(ax.first_col_idx, 11)
        hist.value = "Historical"
        style.label(hist, italic=True)
        if ax.n_forecast:
            fc = self._cell(ax.col_idx(ax.first_forecast), 11)
            fc.value = "Forecast"
            style.label(fc, italic=True)
        # units + year band row 12
        uc = self._cell(style.COL_LABEL, 12)
        uc.value = self.units
        style.band(uc)
        for i in range(ax.n):
            cell = self._cell(ax.col_idx(i), 12)
            cell.value = ax.years[i]
            style.band(cell, year_fmt=ax.year_fmt(i))

    def _r_section(self, s: dict) -> None:
        row = s["row"]
        style.marker(self._cell(style.COL_MARK, row))
        tc = self._cell(style.COL_LABEL, row)
        tc.value = s["title"]
        style.section(tc)
        if s["years"]:
            for i in range(self.axis.n):
                cell = self._cell(self.axis.col_idx(i), row)
                cell.value = self.axis.years[i]
                style.band(cell, year_fmt=self.axis.year_fmt(i))

    def _r_line(self, s: dict) -> None:
        row = s["row"]
        lc = self._cell(style.COL_LABEL, row)
        lc.value = s["label"] or ""
        if s["summary"]:
            style.summary(lc, left_edge=True)
        else:
            style.label(lc, bold=s["bold"], italic=s["italic"])
            if s["total"]:
                lc.border = style.TOP_BORDER
        gc = None
        if s["g"] is not None:
            gc = self._cell(style.COL_INPUT, row)
            gv = s["g"]
            # A MONEY g-cell (currency=True: entry balances) must be FX-converted like every other money
            # input, or flipping the currency toggle mixes source-currency seeds with converted flows and the
            # roll stops tying. It uses the FX rate of `g_fx_col` (default the first year; entry balances at
            # close pass the first-forecast year, so the entry-year roll is internally consistent in either
            # currency). Ratio g-cells (currency=False: multiples, %s) stay literal (FX-invariant).
            if self._fx_row_name and s["currency"] and isinstance(gv, (int, float)):
                fx = f"{self.axis.col(s['g_fx_col'])}${self.rows[self._fx_row_name]}"
                if self._fx_curr_name:
                    gc.value = f"={gv}{self._fx_op}CHOOSE($G${self.rows[self._fx_curr_name]},1,{fx})"
                else:
                    gc.value = f"={gv}{self._fx_op}{fx}"
            else:
                gc.value = gv
            style.input_value(gc, s["g_fmt"], assumption=True)
        if s.get("comment"):
            target = gc if gc is not None else lc
            c = Comment(s["comment"], "model-x-math")
            c.width, c.height = 300, 130
            target.comment = c
        cols = s["cols"] if s["cols"] is not None else range(self.axis.n)
        vals = s["values"]
        for i in cols:
            if i < 0 or i >= self.axis.n:
                continue
            col = self.axis.col(i)
            cell = self._cell(self.axis.col_idx(i), row)
            # A row may MIX inputs (historical actuals) and formulas (forecast):
            # a blue input wins where a value is supplied, else the fn fills in.
            v = vals[i] if (vals is not None and i < len(vals)) else None
            if v is not None:
                # Currency toggle: a money input becomes
                # `=<native>/CHOOSE($G$<curr>, 1, <FX cell>)` — toggle 1 shows the ORIGINAL
                # source-currency number, 2 converts by that year's rate. Editing the toggle
                # or the rate re-drives the whole model. Ratios / FX row (currency=False) stay literal.
                if self._fx_row_name and s["currency"] and isinstance(v, (int, float)):
                    fx = f"{col}${self.rows[self._fx_row_name]}"
                    if self._fx_curr_name:
                        curr = f"$G${self.rows[self._fx_curr_name]}"
                        cell.value = f"={v}{self._fx_op}CHOOSE({curr},1,{fx})"
                    else:
                        cell.value = f"={v}{self._fx_op}{fx}"
                else:
                    cell.value = v
                style.input_value(cell, s["fmt"], bold=s["bold"], assumption=s["assumption"], total=s["total"])
                continue
            if s["fn"] is not None:
                f = s["fn"](self, i, col)
                if f is None:
                    continue
                expr = str(f)
                expr = expr[1:] if expr.startswith("=") else expr
                # Graceful degradation for THIN inputs: ratio rows (%, ×) blank out on a 0/blank
                # denominator instead of throwing #DIV/0!. Value rows stay unguarded so genuine
                # wiring errors still surface to `mxm.verify`. Summary rows (IRR/MoM) manage their
                # own IFERROR, so they're excluded.
                if not s["summary"] and s["fmt"] in style.RATIO_FORMATS and "IFERROR(" not in expr.upper():
                    expr = f'IFERROR({expr},"")'
                cell.value = "=" + expr
                if s["summary"]:
                    style.summary(cell, s["fmt"])
                else:
                    style.value(cell, s["fmt"], bold=s["bold"], italic=s["italic"], total=s["total"])
                    if "!" in expr:  # cross-sheet reference → green font (link colour)
                        cell.font = style.font(bold=s["bold"], italic=s["italic"], color=style.LINK_FONT)

    def _widths(self) -> None:
        from openpyxl.utils import get_column_letter

        for letter, w in style.COL_WIDTHS.items():
            self.ws.column_dimensions[letter].width = w
        for i in range(self.axis.n):
            self.ws.column_dimensions[get_column_letter(self.axis.col_idx(i))].width = style.DATA_WIDTH
