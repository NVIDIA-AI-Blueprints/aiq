"""
Model X Math — house-style tokens.

Every visual constant lives here so the DESIGN is stable across every model the
skill produces. Values were extracted verbatim from the two gold-standard files
(the house-standard "Model X Math" sheets) and the workbook theme, then the
theme tints were resolved to explicit RGB so output does not depend on the
default openpyxl theme.

Do not scatter formatting through blocks.py — call the appliers here.
"""

from __future__ import annotations

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side

# ---- palette (explicit RGB, theme tints pre-resolved) -----------------------
TEXT_DK = "323E48"  # default near-black body text (theme dk1)
HEADER_BAND = "17375E"  # dark navy band: units line + YEAR row (accent1 @ -0.25 tint)
SECTION_BAND = "D9D8D6"  # light-grey section-header band (theme dk2)
SUMMARY_BAND = "F2F2F2"  # very-light band for IRR / MoM summary rows (white @ -0.05)
INPUT_FONT = "0000F4"  # blue — every hardcoded input
LINK_FONT = "008000"  # green — a formula that references another sheet (cross-sheet link)
INPUT_FILL = "FFFFCC"  # light-yellow box behind assumption inputs
WHITE = "FFFFFF"

FONT_NAME = "Calibri"
FONT_SIZE = 11

# ---- number formats (exact strings from the gold standard) ------------------
NUM1 = r"#,##0.0;\(#,##0.0\);\-\-"  # value, 1 decimal
NUM0 = r"#,##0;\(#,##0\);\-\-"  # value, 0 decimals
MULT = r"#,##0.0\x_);\(#,##0.0\x\);\-\-_)"  # multiple, trailing x
PCT = r"0.0%_);\(0.0%\);0.0%_);@_)"  # percent, 1 decimal
IRR_PCT = "0.0%"  # IRR summary cells
YEAR_A = '#"A"'  # actual year (2025A)
YEAR_E = '#"E"'  # forecast/estimate year (2026E)
FX_RATE = "#,##0.00##"  # FX rate cell (90.00, 0.7912) — currency conversion
DATE = "dd/mm/yyyy"

# Ratio-style formats (percent / multiple). Formula rows using these are wrapped in IFERROR by the
# renderer so a thin input with a 0/blank denominator blanks out instead of throwing #DIV/0!.
RATIO_FORMATS = frozenset({PCT, MULT})

# ---- reusable Side/Border objects -------------------------------------------
_THIN = Side(style="thin", color=TEXT_DK)
_HAIR = Side(style="hair", color=TEXT_DK)
TOP_BORDER = Border(top=_THIN)
BOTTOM_BORDER = Border(bottom=_THIN)
BOX_HAIR = Border(top=_HAIR, bottom=_HAIR, left=_HAIR, right=_HAIR)
SUMMARY_BORDER = Border(top=_THIN, left=_THIN)


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


def font(bold: bool = False, italic: bool = False, color: str = TEXT_DK) -> Font:
    return Font(name=FONT_NAME, size=FONT_SIZE, bold=bold, italic=italic, color=color)


RIGHT = Alignment(horizontal="right")
LEFT = Alignment(horizontal="left")
CENTER = Alignment(horizontal="center")

# ---- column geometry (matches the reference models) ------------------------------------------
COL_MARGIN = 1  # A — left margin
COL_MARK = 2  # B — the "x" section marker
COL_LABEL = 3  # C — row labels
COL_INPUT = 7  # G — single-scalar assumption inputs (tax %, rate, split)
COL_DATA0 = 9  # I — first period column

COL_WIDTHS = {"A": 4.5, "B": 2.5, "C": 32.0, "D": 2.0, "E": 2.0, "F": 2.0, "G": 9.0, "H": 2.0}
DATA_WIDTH = 12.5


# ---- appliers ----------------------------------------------------------------
def title(cell: Cell) -> None:
    cell.font = font(bold=True)
    cell.border = BOTTOM_BORDER


def header_label(cell: Cell) -> None:
    """Italic labels like 'Start date' / 'End date'."""
    cell.font = font(italic=True)


def band(cell: Cell, year_fmt: str | None = None) -> None:
    """Dark navy header band (units line + YEAR row): white bold on navy."""
    cell.font = font(bold=True, color=WHITE)
    cell.fill = _fill(HEADER_BAND)
    cell.alignment = RIGHT if year_fmt else LEFT
    if year_fmt:
        cell.number_format = year_fmt


def section(cell: Cell) -> None:
    """Grey section-header band with a bold title."""
    cell.font = font(bold=True)
    cell.fill = _fill(SECTION_BAND)


def marker(cell: Cell) -> None:
    cell.value = "x"
    cell.font = font()


def label(cell: Cell, *, bold: bool = False, italic: bool = False) -> None:
    cell.font = font(bold=bold, italic=italic)


def value(cell: Cell, fmt: str = NUM1, *, bold: bool = False, italic: bool = False, total: bool = False) -> None:
    """A computed (black formula) cell."""
    cell.font = font(bold=bold, italic=italic)
    cell.number_format = fmt
    cell.alignment = RIGHT
    if total:
        cell.border = TOP_BORDER


def input_value(
    cell: Cell, fmt: str = NUM1, *, bold: bool = False, assumption: bool = False, total: bool = False
) -> None:
    """A blue hardcoded input. `assumption=True` adds the yellow hair-box."""
    cell.font = font(bold=bold, italic=assumption, color=INPUT_FONT)
    cell.number_format = fmt
    cell.alignment = CENTER if assumption else RIGHT
    if assumption:
        cell.fill = _fill(INPUT_FILL)
        cell.border = BOX_HAIR
    elif total:
        cell.border = TOP_BORDER


def summary(cell: Cell, fmt: str = NUM1, *, left_edge: bool = False) -> None:
    """A cell in the IRR/MoM summary band."""
    cell.font = font(bold=True)
    cell.fill = _fill(SUMMARY_BAND)
    cell.number_format = fmt
    cell.alignment = RIGHT
    cell.border = SUMMARY_BORDER if left_edge else TOP_BORDER
